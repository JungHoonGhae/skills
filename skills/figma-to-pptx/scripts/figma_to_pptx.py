#!/usr/bin/env python3
"""Figma frames -> fully editable PowerPoint (.pptx).

Boxes, borders and shadows become PowerPoint shapes, vector icons become freeform shapes, text becomes text
boxes in the same fonts as Figma (embedded in the file), image fills become pictures, and Figma layers become
groups. What PowerPoint cannot draw natively (masks, blurs, rotation, blend modes) is placed as a picture of just
that layer and listed in the report.

  FIGMA_TOKEN=... uv run --with python-pptx --with pillow --with fonttools python figma_to_pptx.py "<figma url>" -o out.pptx
  python figma_to_pptx.py --self-test        # offline check of the whole pipeline (no token, no network)

The URL's node-id may point at a page, a section or a single frame. See SKILL.md for every option.
"""
import argparse, collections, hashlib, io, itertools, json, math, os, re, struct, sys, time, urllib.error, urllib.parse, urllib.request
from xml.sax.saxutils import escape

from fontTools.ttLib import TTFont
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu

PX = 6350   # EMU per Figma px: a 1920px frame is a 13.333in slide and 1px of text is 0.5pt
MAX_EMU = 51206400   # PowerPoint's largest slide side (56in)
NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
NOLINE = '<a:ln><a:noFill/></a:ln>'
SLIDES = ('FRAME', 'COMPONENT', 'INSTANCE')
VECTORS = ('VECTOR', 'BOOLEAN_OPERATION', 'STAR', 'LINE', 'REGULAR_POLYGON')
SIDES = ('top', 'right', 'bottom', 'left')
WEIGHTS = {100: 'Thin', 200: 'ExtraLight', 300: 'Light', 400: 'Regular', 500: 'Medium', 600: 'SemiBold', 700: 'Bold', 800: 'ExtraBold', 900: 'Black'}
PANOSE = ('bFamilyType', 'bSerifStyle', 'bWeight', 'bProportion', 'bContrast', 'bStrokeVariation', 'bArmStyle', 'bLetterForm', 'bMidline', 'bXHeight')
BASE = ((0x20, 0x7E), (0xA0, 0xFF), (0x2000, 0x206F), (0x2190, 0x21FF), (0x2460, 0x24FF), (0x25A0, 0x25FF), (0x3000, 0x303F))
HANGUL = ((0x1100, 0x11FF), (0x3130, 0x318F), (0xAC00, 0xD7A3))   # whole Hangul so people can type new Korean text
TOKENS = re.compile(f'[{chr(0xAC00)}-{chr(0xD7A3)}]|[^\\s{chr(0xAC00)}-{chr(0xD7A3)}]+|\\s+')   # Hangul breaks per syllable like Figma
HAN = re.compile(f'[{chr(0xAC00)}-{chr(0xD7A3)}]')
CACHE = os.path.expanduser(os.environ.get('FIGMA_PPTX_CACHE', '~/.cache/figma-to-pptx'))
TOKEN_FILE = os.path.join(os.path.expanduser('~'), '.config', 'figma-to-pptx', 'token')   # lets people save the token without pasting it into a chat
S, DX, DY = 1.0, 0, 0   # ponytail: per-frame scale/offset as globals; the CLI converts one frame at a time
report = collections.Counter()


# ---- Figma API ------------------------------------------------------------------------------------

def api(path):
    token = os.environ.get('FIGMA_TOKEN') or os.environ.get('FIGMA_ACCESS_TOKEN')
    if not token and os.path.exists(TOKEN_FILE):
        token = open(TOKEN_FILE, encoding='utf-8-sig', errors='ignore').read().strip()
    if not token:
        sys.exit(f'No Figma token: set FIGMA_TOKEN or save it in {TOKEN_FILE} '
                 '(Figma > Settings > Security > Personal access tokens, scope File content: read-only). See SKILL.md.')
    for wait in (0, 10, 30, 60, 120):
        time.sleep(wait)
        try:
            req = urllib.request.Request('https://api.figma.com/v1' + path, headers={'X-Figma-Token': token})
            return json.load(urllib.request.urlopen(req, timeout=300))
        except urllib.error.HTTPError as e:
            if e.code != 429:
                sys.exit(f'Figma API {e.code} for {path.split("?")[0]}: {e.read()[:200].decode("utf-8", "replace")}')
    sys.exit('Figma rate limit: try again in a few minutes')


def download(url, path=None):
    if path and os.path.exists(path):
        return open(path, 'rb').read()
    data = urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': 'figma-to-pptx'}), timeout=300).read()
    if path:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        open(path, 'wb').write(data)
    return data


def parse_target(arg, node=None):
    m = re.search(r'figma\.com/(?:design|file|proto|board|slides)/([A-Za-z0-9]+)', arg)
    query = urllib.parse.parse_qs(urllib.parse.urlparse(arg).query) if m else {}
    node = node or (query.get('node-id') or [None])[0]
    return (m.group(1) if m else arg), (node.replace('-', ':') if node else None)


def natural(name):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r'(\d+)', name)]


def reading_order(frames):
    """Rows top to bottom, then left to right inside a row (how people lay slides out on a canvas)."""
    rows = []
    for f in sorted(frames, key=lambda f: (f['absoluteBoundingBox']['y'], f['absoluteBoundingBox']['x'])):
        if rows and abs(f['absoluteBoundingBox']['y'] - rows[-1][0]) < 10:
            rows[-1][1].append(f)
        else:
            rows.append((f['absoluteBoundingBox']['y'], [f]))
    return [f for _, row in rows for f in sorted(row, key=lambda f: f['absoluteBoundingBox']['x'])]


def pick_frames(key, node, match, sort):
    if not node:
        node = api(f'/files/{key}?depth=1')['document']['children'][0]['id']   # no node-id: the first page
    root = api(f'/files/{key}/nodes?ids={node}&depth=2')['nodes'][node]['document']
    found = []
    if root['type'] in ('CANVAS', 'SECTION'):
        for c in root.get('children', []):
            found += [g for g in c.get('children', []) if g['type'] in SLIDES] if c['type'] == 'SECTION' else [c] if c['type'] in SLIDES else []
    elif root['type'] in SLIDES:
        found = [root]
    found = [f for f in found if f.get('visible', True) and f.get('absoluteBoundingBox') and (not match or re.search(match, f['name']))]
    if not found:
        sys.exit(f'no frames under node {node} (type {root["type"]}){" matching " + match if match else ""}')
    return root['name'], (sorted(found, key=lambda f: natural(f['name'])) if sort == 'name' else reading_order(found))


def frame_tree(key, fid, cache):
    path = cache and os.path.join(cache, f'{key}-{fid.replace(":", "-")}.json')
    if path and os.path.exists(path):
        return json.load(open(path))
    doc = api(f'/files/{key}/nodes?ids={fid}&geometry=paths')['nodes'][fid]['document']
    if path:
        os.makedirs(cache, exist_ok=True)
        json.dump(doc, open(path, 'w'))
    return doc


def walk(n):
    if n.get('visible') is False:
        return
    yield n
    for c in n.get('children', []):
        yield from walk(c)


def raster_reason(n):
    """Why this layer must become a picture, or None when PowerPoint can draw it natively."""
    if abs(n.get('rotation') or 0) > 0.01:
        return 'rotated'
    if any(c.get('isMask') and c.get('visible', True) for c in n.get('children', [])):
        return 'mask'
    if any(e.get('visible', True) and e['type'] in ('LAYER_BLUR', 'BACKGROUND_BLUR') for e in n.get('effects') or []):
        return 'blur'
    if n.get('blendMode') not in (None, 'PASS_THROUGH', 'NORMAL'):
        return f'blend {n["blendMode"].lower()}'
    if n['type'] in VECTORS + ('TEXT',) and any(p.get('visible', True) and p['type'] in ('IMAGE', 'VIDEO') for p in n.get('fills') or []):
        return 'image inside a vector or text'
    return None


# ---- fonts ----------------------------------------------------------------------------------------

def bucket(weight):
    return min(900, max(100, int(round((weight or 400) / 100.0)) * 100))


def google_sources(family):
    """Font files of a Google Fonts family from github.com/google/fonts (variable files first)."""
    slug = re.sub(r'[^a-z0-9]', '', family.lower())
    for lic in ('ofl', 'apache', 'ufl'):
        try:
            listing = json.load(urllib.request.urlopen(urllib.request.Request(
                f'https://api.github.com/repos/google/fonts/contents/{lic}/{slug}', headers={'User-Agent': 'figma-to-pptx'}), timeout=60))
        except urllib.error.HTTPError:
            continue
        files = [f['name'] for f in listing if f['name'].endswith('.ttf') and 'Italic' not in f['name']]
        files = [f for f in files if '[' in f] or files
        return [download(f'https://raw.githubusercontent.com/google/fonts/main/{lic}/{slug}/{urllib.parse.quote(f)}',
                         os.path.join(CACHE, 'src', slug, f)) for f in files]
    return []


class Fonts:
    """Figma family + weight -> PowerPoint font name, bold flag and a measuring font; builds the faces to embed.
    PowerPoint only knows regular/bold per family, so other weights become their own family ("Inter Medium")."""

    def __init__(self, overrides, chars, full):
        self.overrides, self.chars, self.full = overrides, chars, full
        self.sources, self.faces, self.embeds, self.ea = {}, {}, collections.defaultdict(dict), None

    def source(self, family):
        if family not in self.sources:
            blobs = [open(p, 'rb').read() for p in self.overrides.get(family, [])] or google_sources(family)
            self.sources[family] = [TTFont(io.BytesIO(b)) for b in blobs]
            if not blobs:
                report[f'font not found, referenced by name only: {family}'] += 1
        return self.sources[family]

    def face(self, family, weight):
        w = bucket(weight)
        if (family, w) in self.faces:
            return self.faces[family, w]
        name = family if w in (400, 700) else f'{family} {WEIGHTS[w]}'
        path = self.build(family, w, name)
        measure = ImageFont.truetype(io.BytesIO(open(path, 'rb').read()), 100) if path else None   # bytes: Windows user folders with Korean names
        if path and self.embeddable(path):
            self.embeds[name]['bold' if w == 700 else 'regular'] = path
        self.faces[family, w] = (name, w == 700, measure)
        return self.faces[family, w]

    def build(self, family, w, name):
        """Static TrueType face for this weight, cut to the characters in the deck (plus all Hangul when Korean
        is used), with its internal names rewritten to the PowerPoint font name (PowerPoint ignores it otherwise)."""
        fonts = self.source(family)
        if not fonts:
            return None
        from fontTools import subset
        from fontTools.varLib import instancer
        var = next((f for f in fonts if 'fvar' in f), None)
        static = min(fonts, key=lambda f: abs(f['OS/2'].usWeightClass - w)) if not var else None
        keep = set(self.chars.get(family, ())) | {u for a, b in BASE for u in range(a, b + 1)}
        if any(HANGUL[2][0] <= u <= HANGUL[2][1] for u in keep):
            keep |= {u for a, b in HANGUL for u in range(a, b + 1)}
        sig = hashlib.sha1(json.dumps([family, w, self.full, sorted(keep)]).encode()).hexdigest()[:12]
        path = os.path.join(CACHE, 'faces', f'{re.sub(r"[^A-Za-z0-9]", "", name)}-{sig}.ttf')
        if os.path.exists(path):
            return path
        f = TTFont(io.BytesIO(_bytes(var or static)))
        if not self.full:
            opt = subset.Options()
            opt.layout_features, opt.name_IDs, opt.name_languages, opt.notdef_outline = ['*'], ['*'], ['*'], True
            sub = subset.Subsetter(opt)
            sub.populate(unicodes=sorted(keep))
            sub.subset(f)
        if 'fvar' in f:
            axes = {a.axisTag: a.defaultValue for a in f['fvar'].axes}
            axes['wght'] = max(min(w, max(a.maxValue for a in f['fvar'].axes if a.axisTag == 'wght')),
                               min(a.minValue for a in f['fvar'].axes if a.axisTag == 'wght')) if 'wght' in axes else None
            f = instancer.instantiateVariableFont(f, {k: v for k, v in axes.items() if v is not None})
        style = 'Bold' if w == 700 else 'Regular'
        f['name'].names = [r for r in f['name'].names if r.nameID not in (1, 2, 4, 6, 16, 17)]
        for i, v in {1: name, 2: style, 4: f'{name} {style}' if w == 700 else name, 6: f'{re.sub(r"[^A-Za-z0-9]", "", name)}-{style}'}.items():
            f['name'].setName(v, i, 3, 1, 0x409)
        f['OS/2'].usWeightClass = w
        f['OS/2'].fsSelection = (f['OS/2'].fsSelection & ~0b1100001) | (0b100000 if w == 700 else 0b1000000)
        f['head'].macStyle = 1 if w == 700 else 0
        os.makedirs(os.path.dirname(path), exist_ok=True)
        f.save(path)
        return path

    def embeddable(self, path):
        f = TTFont(path)
        if 'glyf' not in f:
            report['font is OpenType CFF, PowerPoint can only embed TrueType (not embedded): ' + f['name'].getDebugName(1)] += 1
            return False
        if f['OS/2'].fsType & 0x0202:   # restricted license or bitmap-only embedding
            report['font license forbids embedding (not embedded): ' + f['name'].getDebugName(1)] += 1
            return False
        return True

    def has(self, family, cp):
        return any(cp in (f.getBestCmap() or {}) for f in self.source(family))


def _bytes(font):
    buf = io.BytesIO()
    font.save(buf)
    return buf.getvalue()


def eot(path):
    """TrueType -> uncompressed EOT (version 0x00020001, flags 0): the only .fntdata payload PowerPoint reads."""
    data, f = open(path, 'rb').read(), TTFont(path)
    o, nm = f['OS/2'], f['name']
    head = struct.pack('<LL10sBBLHH6LL4L', 0x00020001, 0, bytes(getattr(o.panose, k) for k in PANOSE), 1, 0, o.usWeightClass,
                       o.fsType, 0x504C, o.ulUnicodeRange1, o.ulUnicodeRange2, o.ulUnicodeRange3, o.ulUnicodeRange4,
                       getattr(o, 'ulCodePageRange1', 0), getattr(o, 'ulCodePageRange2', 0), f['head'].checkSumAdjustment, 0, 0, 0, 0)
    for i in (1, 2, 5, 4):   # family, style, version, full name (UTF-16LE)
        s = (nm.getDebugName(i) or '').encode('utf-16-le')
        head += struct.pack('<HH', 0, len(s)) + s
    head += struct.pack('<HH', 0, 0)   # no RootString
    return struct.pack('<LL', 8 + len(head) + len(data), len(data)) + head + data


def embed(prs, fonts):
    """Put the faces in ppt/fonts and register them; PowerPoint ignores them unless embedTrueTypeFonts="1"."""
    items, n = '', 0
    for name, slots in sorted(fonts.embeds.items()):
        refs = ''
        for slot in ('regular', 'bold'):
            if slot in slots:
                n += 1
                part = Part(PackURI(f'/ppt/fonts/font{n}.fntdata'), 'application/x-fontdata', prs.part.package, eot(slots[slot]))
                refs += f'<p:{slot} r:id="{prs.part.relate_to(part, RT.FONT)}"/>'
        f = TTFont(next(iter(slots.values())))
        cmap = f.getBestCmap() or {}
        charset = '-127' if 0xAC00 in cmap else '-128' if 0x3042 in cmap else '-122' if 0x4E00 in cmap else '0'
        items += (f'<p:embeddedFont><p:font typeface="{escape(name)}" panose="{"".join("%02X" % getattr(f["OS/2"].panose, k) for k in PANOSE)}" '
                  f'pitchFamily="34" charset="{charset}"/>{refs}</p:embeddedFont>')
    if items:
        prs.part._element.set('embedTrueTypeFonts', '1')
        prs.part._element.find(qn('p:notesSz')).addnext(parse_xml(f'<p:embeddedFontLst {NS}>{items}</p:embeddedFontLst>'))


# ---- shapes ---------------------------------------------------------------------------------------

def emu(v):
    return round(v * PX * S)


def rgb(c):
    return '%02X%02X%02X' % tuple(round(c[k] * 255) for k in 'rgb')


def paint(paints, alpha=1.0):
    """Topmost visible solid or gradient paint -> ('solid', hex, alpha) or ('grad', kind, stops, handles)."""
    for p in reversed(paints or []):
        if not p.get('visible', True) or p['type'] in ('IMAGE', 'VIDEO'):
            continue
        a = p.get('opacity', 1) * alpha
        if p['type'] == 'SOLID':
            return ('solid', rgb(p['color']), p['color'].get('a', 1) * a)
        if p['type'].startswith('GRADIENT'):
            return ('grad', p['type'], [(s['position'], rgb(s['color']), s['color'].get('a', 1) * a) for s in p['gradientStops']],
                    p.get('gradientHandlePositions'))
    return None


def color_xml(hexc, a):
    return f'<a:srgbClr val="{hexc}">{f"<a:alpha val=\"{round(a * 100000)}\"/>" if a < 0.999 else ""}</a:srgbClr>'


def fill_xml(c, w=1, h=1):
    if not c:
        return '<a:noFill/>'
    if c[0] == 'solid':
        return f'<a:solidFill>{color_xml(c[1], c[2])}</a:solidFill>'
    _, kind, stops, hp = c
    gs = ''.join(f'<a:gs pos="{round(min(1, max(0, p)) * 100000)}">{color_xml(hx, a)}</a:gs>' for p, hx, a in stops)
    if kind == 'GRADIENT_LINEAR' and hp:
        ang = math.degrees(math.atan2((hp[1]['y'] - hp[0]['y']) * h, (hp[1]['x'] - hp[0]['x']) * w)) % 360
        shade = f'<a:lin ang="{round(ang * 60000)}" scaled="0"/>'
    else:
        if kind != 'GRADIENT_RADIAL':
            report[f'{kind.lower()} drawn as radial'] += 1
        shade = '<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>'
    return f'<a:gradFill rotWithShape="1"><a:gsLst>{gs}</a:gsLst>{shade}</a:gradFill>'


def solid_of(c):
    """Single colour for places that cannot take a gradient (text, borders): the first stop."""
    return c if not c or c[0] == 'solid' else ('solid', c[2][0][1], c[2][0][2])


def prst(kind, adj=()):
    gd = ''.join(f'<a:gd name="adj{i + 1 if len(adj) > 1 else ""}" fmla="val {v}"/>' for i, v in enumerate(adj))
    return f'<a:prstGeom prst="{kind}"><a:avLst>{gd}</a:avLst></a:prstGeom>'


def corners(n, w, h, inset):
    """Corner radii -> preset: equal corners -> rounded rectangle, equal top and bottom pairs -> round2SameRect."""
    tl, tr, br, bl = n.get('rectangleCornerRadii') or [n.get('cornerRadius', 0) or 0] * 4
    ratio = lambda r: min(50000, round(max(0, r - inset) / max(min(w, h), 1e-6) * 100000))
    if tl == tr == br == bl:
        return prst('roundRect', [ratio(tl)]) if tl else prst('rect')
    if tl == tr and bl == br:
        return prst('round2SameRect', [ratio(tl), ratio(bl)])
    report['mixed corner radii drawn with the largest radius'] += 1
    return prst('roundRect', [ratio(max(tl, tr, br, bl))])


def effects(n, alpha):
    out = {}
    for e in n.get('effects') or []:
        if not e.get('visible', True) or e['type'] not in ('DROP_SHADOW', 'INNER_SHADOW'):
            continue
        tag = 'outerShdw' if e['type'] == 'DROP_SHADOW' else 'innerShdw'
        if tag in out:
            report['extra shadows dropped (PowerPoint keeps one of each kind)'] += 1
            continue
        ox, oy = e['offset']['x'], e['offset']['y']
        extra = ' algn="ctr" rotWithShape="0"' if tag == 'outerShdw' else ''
        out[tag] = (f'<a:{tag} blurRad="{emu(e.get("radius", 0))}" dist="{emu(math.hypot(ox, oy))}" '
                    f'dir="{round(math.degrees(math.atan2(oy, ox)) % 360 * 60000)}"{extra}>{color_xml(rgb(e["color"]), e["color"].get("a", 1) * alpha)}</a:{tag}>')
    return f'<a:effectLst>{out.get("innerShdw", "")}{out.get("outerShdw", "")}</a:effectLst>' if out else ''


def esc(name):
    return escape(name[:80], {'"': '&quot;'})


class Sheet:
    """One slide. Shapes are written as DrawingML directly (freeforms, shadows, transparency, groups)."""

    def __init__(self, slide, ox, oy):
        self.slide, self.stack, self.n, self.ox, self.oy = slide, [[slide.shapes._spTree, None]], 1, ox, oy

    def put(self, el, x, y, w, h):
        self.stack[-1][0].append(el)
        for g in self.stack[1:]:
            b = g[1] or (x, y, x + w, y + h)
            g[1] = (min(b[0], x), min(b[1], y), max(b[2], x + w), max(b[3], y + h))

    def xfrm(self, x, y, w, h):
        return f'<a:xfrm><a:off x="{emu(x) + DX}" y="{emu(y) + DY}"/><a:ext cx="{max(1, emu(w))}" cy="{max(1, emu(h))}"/></a:xfrm>'

    def add(self, name, x, y, w, h, geom, fill=None, line=NOLINE, effect='', text=''):
        self.n += 1
        nv = '<p:cNvSpPr txBox="1"/>' if text else '<p:cNvSpPr/>'
        self.put(parse_xml(f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{self.n}" name="{esc(name)}"/>{nv}<p:nvPr/></p:nvSpPr>'
                           f'<p:spPr>{self.xfrm(x, y, w, h)}{geom}{fill_xml(fill, w, h)}{line}{effect}</p:spPr>{text}</p:sp>'), x, y, w, h)
        report['text boxes' if text else 'shapes'] += 1

    def picture(self, name, x, y, w, h, png, geom=None, crop=None):
        rid = self.slide.part.get_or_add_image_part(io.BytesIO(png))[1]
        src = f'<a:srcRect l="{crop[0]}" t="{crop[1]}" r="{crop[2]}" b="{crop[3]}"/>' if crop else ''
        self.n += 1
        self.put(parse_xml(f'<p:pic {NS}><p:nvPicPr><p:cNvPr id="{self.n}" name="{esc(name)}"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
                           f'<p:blipFill><a:blip r:embed="{rid}"/>{src}<a:stretch><a:fillRect/></a:stretch></p:blipFill>'
                           f'<p:spPr>{self.xfrm(x, y, w, h)}{geom or prst("rect")}</p:spPr></p:pic>'), x, y, w, h)
        report['pictures'] += 1

    def path(self, name, x, y, d, color):
        """SVG path in node coordinates -> freeform shape placed at the path's own bounds."""
        toks, cmds, pts, i = re.findall(r'[MLCQZ]|-?\d*\.?\d+(?:e-?\d+)?', d), [], [], 0
        while i < len(toks):
            c = toks[i]
            k = {'M': 2, 'L': 2, 'C': 6, 'Q': 4, 'Z': 0}[c]
            v = [float(t) for t in toks[i + 1:i + 1 + k]]
            cmds.append((c, v))
            pts += zip(v[0::2], v[1::2])
            i += 1 + k
        if not pts:
            return
        x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
        w, h = max(p[0] for p in pts) - x0, max(p[1] for p in pts) - y0
        pt = lambda px, py: f'<a:pt x="{emu(px - x0)}" y="{emu(py - y0)}"/>'
        tag = {'M': 'moveTo', 'L': 'lnTo', 'C': 'cubicBezTo', 'Q': 'quadBezTo'}
        body = ''.join('<a:close/>' if c == 'Z' else f'<a:{tag[c]}>{"".join(pt(*v[j:j + 2]) for j in range(0, len(v), 2))}</a:{tag[c]}>' for c, v in cmds)
        self.add(name, x + x0, y + y0, w, h,
                 f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="l" t="t" r="r" b="b"/>'
                 f'<a:pathLst><a:path w="{max(1, emu(w))}" h="{max(1, emu(h))}">{body}</a:path></a:pathLst></a:custGeom>', color)

    def open(self, name):
        self.n += 1
        g = parse_xml(f'<p:grpSp {NS}><p:nvGrpSpPr><p:cNvPr id="{self.n}" name="{esc(name)}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
                      f'<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:grpSp>')
        self.stack[-1][0].append(g)
        self.stack.append([g, None])

    def close(self):
        g, b = self.stack.pop()
        kids = [c for c in g if c.tag in (qn('p:sp'), qn('p:grpSp'), qn('p:pic'))]
        if len(kids) < 2:   # a group of one is just noise in the selection pane
            for c in kids:
                g.addprevious(c)
            g.getparent().remove(g)
            return
        xf = g.find(qn('p:grpSpPr')).find(qn('a:xfrm'))
        x, y, cx, cy = emu(b[0]) + DX, emu(b[1]) + DY, max(1, emu(b[2] - b[0])), max(1, emu(b[3] - b[1]))
        for tag, a, v in (('a:off', 'x', x), ('a:off', 'y', y), ('a:ext', 'cx', cx), ('a:ext', 'cy', cy),
                          ('a:chOff', 'x', x), ('a:chOff', 'y', y), ('a:chExt', 'cx', cx), ('a:chExt', 'cy', cy)):
            xf.find(qn(tag)).set(a, str(v))   # child space = slide space, so children keep absolute positions
        report['groups'] += 1


# ---- text -----------------------------------------------------------------------------------------

def paragraphs(n):
    """Text node -> paragraphs of (text, style, fills) runs, split where the style changes."""
    chars = n['characters'].replace(chr(0x2028), '\n')
    ids, table, st = n.get('characterStyleOverrides') or [], n.get('styleOverrideTable') or {}, n['style']
    paras, cur, i = [], [], 0
    while i <= len(chars):
        if i == len(chars) or chars[i] == '\n':
            paras.append(cur)
            cur, i = [], i + 1
            continue
        sid, j = (ids[i] if i < len(ids) else 0), i
        while j < len(chars) and chars[j] != '\n' and (ids[j] if j < len(ids) else 0) == sid:
            j += 1
        over = table.get(str(sid), {})
        cur.append((chars[i:j], {**st, **over}, over.get('fills')))
        i = j
    return paras


class Text:
    def __init__(self, fonts):
        self.fonts = fonts

    def fams(self, s, ch=None):
        """(latin family, east-asian family) for a run; text the font cannot show falls back to the deck's CJK font."""
        fam = s.get('fontFamily') or 'Inter'
        ea = fam if (self.fonts.ea is None or self.fonts.has(fam, 0xAC00) or self.fonts.has(fam, 0x4E00)) else self.fonts.ea
        return fam, ea

    def width(self, t, s):
        total = 0.0
        for latin, seg in itertools.groupby(t, key=lambda ch: ord(ch) < 0x300):
            seg = ''.join(seg)
            fam, ea = self.fams(s)
            m = self.fonts.face(fam if latin else ea, s.get('fontWeight', 400))[2]
            if m is None:
                return None
            total += m.getlength(seg) * s.get('fontSize', 16) / 100
        return total + (s.get('letterSpacing') or 0) * len(t)

    def lines(self, para, w, k):
        n, x, w = 1, 0.0, max(w, 1)
        for t, s, _ in para:
            for tok in TOKENS.findall(t):
                tw = self.width(tok, s)
                if tw is None:
                    return 1
                tw *= k
                if not tok.isspace() and x > 0 and x + tw > w:
                    n, x = n + 1, 0.0
                if tw > w and not tok.isspace():
                    extra = math.ceil(tw / w) - 1
                    n, tw = n + extra, tw - extra * w
                x += tw
        return n

    def add(self, sheet, n, x, y, w, h, alpha):
        st, paras = n['style'], paragraphs(n)
        lh = st.get('lineHeightPx') or st.get('fontSize', 16) * 1.4
        single = st.get('textAutoResize') == 'WIDTH_AND_HEIGHT' or (len(paras) == 1 and h < lh * 1.5)
        types, levels, out = n.get('lineTypes') or [], n.get('lineIndentations') or [], []
        kinds = [types[i] if i < len(types) else 'NONE' for i in range(len(paras))]
        marl = [st.get('fontSize', 16) * 1.1 * max(1, levels[i] if i < len(levels) else 1) if kinds[i] in ('UNORDERED', 'ORDERED') else 0
                for i in range(len(paras))]
        cap = st.get('maxLines') or max(len(paras), round(h / lh))

        def fits(ps, k=1.0):
            if single:
                ws = [sum(self.width(t, s) or 0 for t, s, _ in p) + m for p, m in zip(ps, marl)]
                return max(ws) * k <= w + 1
            return sum(self.lines(p, w - m, k) for p, m in zip(ps, marl)) <= cap

        if st.get('textTruncation') == 'ENDING' and not n.get('characterStyleOverrides') and not fits(paras):
            text, lo, hi = n['characters'].replace(chr(0x2028), '\n'), 0, len(n['characters'])
            while lo < hi:   # longest prefix that still fits with an ellipsis, like Figma's truncation
                mid = (lo + hi + 1) // 2
                lo, hi = (mid, hi) if fits([[(p, st, None)] for p in (text[:mid].rstrip() + '…').split('\n')]) else (lo, mid - 1)
            paras = [[(p, st, None)] for p in (text[:lo].rstrip() + '…').split('\n')]
            marl, kinds = (marl + [0] * len(paras))[:len(paras)], (kinds + ['NONE'] * len(paras))[:len(paras)]
            report['text truncated with an ellipsis'] += 1
        k = 1.0
        while k > 0.8 and any(paras) and not fits(paras, k):   # PowerPoint and Figma wrap a little differently
            k -= 0.02
        if k < 1:
            report['text boxes shrunk slightly to keep Figma line breaks'] += 1
        for i, para in enumerate(paras):
            attrs = f' algn="{ {"CENTER": "ctr", "RIGHT": "r", "JUSTIFIED": "just"}.get(st.get("textAlignHorizontal"), "l")}" eaLnBrk="1" latinLnBrk="0"'
            ppr = f'<a:lnSpc><a:spcPts val="{round(lh * 50 * S)}"/></a:lnSpc>'
            if st.get('paragraphSpacing'):
                ppr += f'<a:spcAft><a:spcPts val="{round(st["paragraphSpacing"] * 50 * S)}"/></a:spcAft>'
            if marl[i]:
                attrs += f' marL="{emu(marl[i])}" indent="{-emu(st.get("fontSize", 16) * 1.1)}"'
                ppr += '<a:buFont typeface="Arial"/><a:buChar char="•"/>' if kinds[i] == 'UNORDERED' else '<a:buAutoNum type="arabicPeriod"/>'
            else:
                ppr += '<a:buNone/>'
            runs = ''
            for t, s, fills in para:
                t = t.upper() if s.get('textCase') == 'UPPER' else t.lower() if s.get('textCase') == 'LOWER' else t
                fam, ea = self.fams(s)
                latin, bold, _ = self.fonts.face(fam, s.get('fontWeight', 400))
                ea_name = self.fonts.face(ea, s.get('fontWeight', 400))[0]
                extra = (' b="1"' if bold else '') + (' i="1"' if s.get('italic') else '')
                extra += f' spc="{round(s["letterSpacing"] * 50 * S)}"' if s.get('letterSpacing') else ''
                extra += {'UNDERLINE': ' u="sng"', 'STRIKETHROUGH': ' strike="sngStrike"'}.get(s.get('textDecoration'), '')
                link = ''
                if (s.get('hyperlink') or {}).get('type') == 'URL':
                    link = f'<a:hlinkClick r:id="{sheet.slide.part.relate_to(s["hyperlink"]["url"], RT.HYPERLINK, is_external=True)}"/>'
                color = solid_of(paint(fills or n.get('fills'), alpha)) or ('solid', '000000', 1)
                keep = ' xml:space="preserve"' if t != t.strip() else ''
                runs += (f'<a:r><a:rPr lang="{"ko-KR" if HAN.search(t) else "en-US"}" '
                         f'sz="{max(100, round(s.get("fontSize", 16) * 50 * k * S))}"{extra} dirty="0">{fill_xml(color)}'
                         f'<a:latin typeface="{escape(latin)}"/><a:ea typeface="{escape(ea_name)}"/>{link}</a:rPr>'
                         f'<a:t{keep}>{escape(re.sub(r"[\x00-\x08\x0b-\x1f]", "", t))}</a:t></a:r>')
            out.append(f'<a:p><a:pPr{attrs}>{ppr}</a:pPr>{runs}<a:endParaRPr sz="{max(100, round(st.get("fontSize", 16) * 50 * k * S))}" dirty="0"/></a:p>')
        anchor = {'CENTER': 'ctr', 'BOTTOM': 'b'}.get(st.get('textAlignVertical'), 't')
        body = (f'<p:txBody><a:bodyPr wrap="{"none" if single else "square"}" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}" rtlCol="0">'
                f'<a:noAutofit/></a:bodyPr><a:lstStyle/>{"".join(out)}</p:txBody>')
        sheet.add(n['name'], x, y, w, h, prst('rect'), text=body)


# ---- drawing --------------------------------------------------------------------------------------

class Painter:
    def __init__(self, text, images, rasters, depth):
        self.text, self.images, self.rasters, self.depth = text, images, rasters, depth

    def image(self, sheet, n, x, y, w, h, geom):
        p = next(p for p in reversed(n['fills']) if p.get('visible', True) and p['type'] in ('IMAGE', 'VIDEO'))
        data = self.images.get(p.get('imageRef'))
        if not data:
            report['image fill missing from Figma (left empty)'] += 1
            return
        iw, ih = Image.open(io.BytesIO(data)).size
        mode, crop = p.get('scaleMode', 'FILL'), None
        if mode == 'FIT':   # contain: shrink the picture box, keep the aspect ratio
            k = min(w / iw, h / ih)
            x, y, w, h = x + (w - iw * k) / 2, y + (h - ih * k) / 2, iw * k, ih * k
        else:   # FILL (cover); CROP and TILE are approximated the same way
            if mode not in ('FILL',):
                report[f'image {mode.lower()} drawn as fill'] += 1
            k = max(w / iw, h / ih)
            cx, cy = (iw - w / k) / 2 / iw, (ih - h / k) / 2 / ih
            crop = tuple(round(v * 100000) for v in (cx, cy, cx, cy))
        sheet.picture(n['name'], x, y, w, h, data, geom, crop)

    def draw(self, sheet, n, clip, alpha=1.0, root=False, depth=0):
        bb = n.get('absoluteBoundingBox')
        if n.get('visible') is False or not bb:
            return
        x, y, w, h = bb['x'] - sheet.ox, bb['y'] - sheet.oy, bb['width'], bb['height']
        if x >= clip[2] or y >= clip[3] or x + w <= clip[0] or y + h <= clip[1]:
            return   # clipped away by a parent (scrolled content and the like)
        alpha *= n.get('opacity', 1)
        if n['id'] in self.rasters and not root:
            rb = n.get('absoluteRenderBounds') or bb
            sheet.picture(n['name'], rb['x'] - sheet.ox, rb['y'] - sheet.oy, rb['width'], rb['height'], self.rasters[n['id']])
            return
        if n['type'] == 'TEXT':
            return self.text.add(sheet, n, x, y, w, h, alpha)
        fill, stroke, sw = paint(n.get('fills'), alpha), solid_of(paint(n.get('strokes'), alpha)), n.get('strokeWeight') or 0
        has_image = any(p.get('visible', True) and p['type'] in ('IMAGE', 'VIDEO') for p in n.get('fills') or [])
        sides = n.get('individualStrokeWeights')
        if sides and len({sides.get(s, 0) for s in SIDES}) == 1:
            sw, sides = sides.get('top', 0), None   # equal sides: one border that follows rounded corners
        arc = n.get('arcData') or {}
        if n['type'] in VECTORS or (n['type'] == 'ELLIPSE' and (arc.get('innerRadius') or abs(arc.get('endingAngle', 2 * math.pi) - arc.get('startingAngle', 0)) < 6.28)):
            for g in (n.get('fillGeometry') or []) if fill else []:
                sheet.path(n['name'], x, y, g['path'], fill)
            for g in (n.get('strokeGeometry') or []) if stroke and sw else []:
                sheet.path(n['name'] + ' stroke', x, y, g['path'], stroke)
            return
        border = bool(stroke and sw and not sides)
        kids = any(c.get('visible') is not False for c in n.get('children', []))
        group = kids and 0 < depth <= self.depth
        if group:
            sheet.open(n['name'])
        inset = {'INSIDE': sw / 2, 'OUTSIDE': -sw / 2}.get(n.get('strokeAlign'), 0) if border else 0
        bx, by = max(x + inset, clip[0]), max(y + inset, clip[1])
        bw, bh = min(x + w - inset, clip[2]) - bx, min(y + h - inset, clip[3]) - by
        cut = (bx, by, bw, bh) != (x + inset, y + inset, w - 2 * inset, h - 2 * inset)
        geom = prst('ellipse') if n['type'] == 'ELLIPSE' else prst('rect') if cut else corners(n, bw, bh, inset)
        dash = '<a:prstDash val="dash"/>' if n.get('strokeDashes') else ''
        line = f'<a:ln w="{emu(sw)}">{fill_xml(stroke)}{dash}</a:ln>' if border else NOLINE
        if bw > 0 and bh > 0 and not (root and not has_image):
            if fill and not root:
                sheet.add(n['name'], bx, by, bw, bh, geom, fill, NOLINE if kids or has_image else line, effects(n, alpha))
            if has_image:
                self.image(sheet, n, bx, by, bw, bh, geom)
            if border and not kids and not fill and not has_image:
                sheet.add(n['name'], bx, by, bw, bh, geom, None, line, effects(n, alpha))
        inner = (max(clip[0], x), max(clip[1], y), min(clip[2], x + w), min(clip[3], y + h)) if n.get('clipsContent') else clip
        for c in n.get('children', []):
            self.draw(sheet, c, inner, alpha, depth=depth + 1)
        if border and (kids or root or has_image) and bw > 0 and bh > 0:   # Figma paints frame borders above content
            sheet.add(n['name'] + ' border', bx, by, bw, bh, geom, None, line)
        if stroke and sides:   # one-sided borders (table rules and the like) as thin bars above content
            t, r, b, l = (sides.get(s, 0) for s in SIDES)
            for rx, ry, rw, rh in ((x, y, w, t), (x + w - r, y, r, h), (x, y + h - b, w, b), (x, y, l, h)):
                if rw > 0 and rh > 0:
                    sheet.add(n['name'] + ' border', rx, ry, rw, rh, prst('rect'), stroke)
        if group:
            sheet.close()


# ---- pipeline -------------------------------------------------------------------------------------

def convert(trees, out, key=None, overrides=None, embed_fonts=True, full=False, depth=3, images=None, rasters=None, link=None):
    global S, DX, DY
    chars, used = collections.defaultdict(set), collections.Counter()
    for t in trees:
        for n in walk(t):
            if n['type'] == 'TEXT':
                for p in paragraphs(n):
                    for seg, s, _ in p:
                        fam = s.get('fontFamily') or 'Inter'
                        chars[fam] |= {ord(c) for c in seg}
                        used[fam] += len(seg)
    fonts = Fonts(overrides or {}, chars, full)
    cjk = [f for f, _ in used.most_common() if any(0xAC00 <= c <= 0xD7A3 or 0x4E00 <= c <= 0x9FFF for c in chars[f]) and (fonts.has(f, 0xAC00) or fonts.has(f, 0x4E00))]
    fonts.ea = cjk[0] if cjk else None
    for fam in list(chars):   # characters a latin-only font cannot show go to the CJK font, so cut them into that one too
        if fonts.ea and fam != fonts.ea:
            chars[fonts.ea] |= {c for c in chars[fam] if c >= 0x300}
    painter = Painter(Text(fonts), images or {}, rasters or {}, depth)
    prs = Presentation()
    fw, fh = trees[0]['absoluteBoundingBox']['width'], trees[0]['absoluteBoundingBox']['height']
    k0 = min(1.0, MAX_EMU / (fw * PX), MAX_EMU / (fh * PX))
    prs.slide_width, prs.slide_height = Emu(round(fw * PX * k0)), Emu(round(fh * PX * k0))
    for t in trees:
        bb = t['absoluteBoundingBox']
        S = min(prs.slide_width / (bb['width'] * PX), prs.slide_height / (bb['height'] * PX))
        DX, DY = round((prs.slide_width - bb['width'] * PX * S) / 2), round((prs.slide_height - bb['height'] * PX * S) / 2)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = solid_of(paint(t.get('fills')))
        if bg:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(bg[1])
        painter.draw(Sheet(slide, bb['x'], bb['y']), t, (0, 0, bb['width'], bb['height']), root=True)
        note = t['name'] + (f'\nhttps://www.figma.com/design/{key}/?node-id={t["id"].replace(":", "-")}' if key else '')
        slide.notes_slide.notes_text_frame.text = note
        print(f'  {len(prs.slides):3d}  {t["name"]}', flush=True)
    if embed_fonts:
        embed(prs, fonts)
    prs.save(out)
    return fonts


def run(args):
    key, node = parse_target(args.target, args.node)
    title, frames = pick_frames(key, node, args.match, args.sort)
    print(f'{len(frames)} frames from "{title}"', flush=True)
    trees = [frame_tree(key, f['id'], args.cache) for f in frames]
    refs = {p['imageRef'] for t in trees for n in walk(t) for p in n.get('fills') or [] if p.get('visible', True) and p.get('imageRef')}
    images = {}
    if refs:
        urls = api(f'/files/{key}/images')['meta']['images']
        images = {r: download(urls[r], os.path.join(CACHE, 'images', r)) for r in refs if urls.get(r)}
    need, seen = {}, set()
    for t in trees:
        for n in walk(t):
            if n['id'] not in seen and n is not t and (reason := raster_reason(n)):
                need[n['id']] = reason
                seen |= {d['id'] for d in walk(n)}   # descendants are inside the picture
    rasters = {}
    ids = sorted(need)
    for i in range(0, len(ids), 40):
        got = api(f'/images/{key}?ids={",".join(ids[i:i + 40])}&format=png&scale=2')['images']
        rasters.update({nid: download(u) for nid, u in got.items() if u})
    for nid in rasters:
        report[f'pictures for layers PowerPoint cannot draw ({need[nid]})'] += 1
    out = args.out or (re.sub(r'[\\/:*?"<>|]+', '_', title).strip(' .')[:100] or 'figma') + '.pptx'
    overrides = {}
    for spec in args.font:
        fam, _, files = spec.partition('=')
        overrides[fam.strip()] = [os.path.expanduser(p.strip()) for p in files.split(',') if p.strip()]
    fonts = convert(trees, out, key, overrides, not args.no_embed, args.full_fonts, args.group_depth, images, rasters)
    print(f'\nsaved {out} ({os.path.getsize(out) / 1e6:.1f} MB)')
    for k, v in sorted(report.items()):
        print(f'  {k}: {v}')
    if fonts.embeds:
        print('  embedded fonts: ' + ', '.join(f'{k} ({"/".join(v)})' for k, v in sorted(fonts.embeds.items())))


def self_test():
    """Offline check: a synthetic frame with a box, a border, an icon, text in a tiny built font and a group."""
    import tempfile
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    global CACHE
    tmp = tempfile.mkdtemp()
    CACHE = os.path.join(tmp, 'cache')
    fb = FontBuilder(1000, isTTF=True)
    fb.setupGlyphOrder(['.notdef', 'A', 'space'])
    fb.setupCharacterMap({65: 'A', 32: 'space'})
    pen = TTGlyphPen(None)
    pen.moveTo((0, 0)), pen.lineTo((500, 0)), pen.lineTo((250, 700)), pen.closePath()
    glyph = pen.glyph()
    fb.setupGlyf({'.notdef': glyph, 'A': glyph, 'space': TTGlyphPen(None).glyph()})
    fb.setupHorizontalMetrics({'.notdef': (500, 0), 'A': (500, 0), 'space': (250, 0)})
    fb.setupHorizontalHeader(ascent=800, descent=-200)
    fb.setupNameTable({'familyName': 'Probe', 'styleName': 'Regular'})
    fb.setupOS2(usWeightClass=400, fsType=0)
    fb.setupPost()
    font = os.path.join(tmp, 'Probe.ttf')
    fb.save(font)
    box = lambda i, x, y, w, h, **kw: {'id': i, 'name': i, 'type': 'RECTANGLE', 'absoluteBoundingBox': {'x': x, 'y': y, 'width': w, 'height': h}, **kw}
    solid = lambda r, g, b: [{'type': 'SOLID', 'color': {'r': r, 'g': g, 'b': b, 'a': 1}}]
    tree = {'id': '1:1', 'name': 'Frame', 'type': 'FRAME', 'absoluteBoundingBox': {'x': 0, 'y': 0, 'width': 400, 'height': 300},
            'fills': solid(1, 1, 1), 'clipsContent': True, 'children': [
                {'id': '1:2', 'name': 'Card', 'type': 'FRAME', 'absoluteBoundingBox': {'x': 20, 'y': 20, 'width': 200, 'height': 100},
                 'fills': solid(0.9, 0.9, 0.9), 'strokes': solid(0, 0, 0), 'strokeWeight': 1, 'strokeAlign': 'INSIDE', 'cornerRadius': 8,
                 'effects': [{'type': 'DROP_SHADOW', 'color': {'r': 0, 'g': 0, 'b': 0, 'a': 0.2}, 'offset': {'x': 0, 'y': 2}, 'radius': 4}],
                 'children': [box('1:3', 30, 30, 50, 20, fills=solid(1, 0, 0)),
                              {'id': '1:4', 'name': 'Label', 'type': 'TEXT', 'characters': 'AA A', 'absoluteBoundingBox': {'x': 30, 'y': 60, 'width': 120, 'height': 20},
                               'style': {'fontFamily': 'Probe', 'fontWeight': 400, 'fontSize': 16, 'lineHeightPx': 20, 'textAutoResize': 'WIDTH_AND_HEIGHT'},
                               'fills': solid(0, 0, 0)}]},
                {'id': '1:5', 'name': 'Icon', 'type': 'VECTOR', 'absoluteBoundingBox': {'x': 300, 'y': 40, 'width': 24, 'height': 24},
                 'fills': solid(0, 0, 1), 'fillGeometry': [{'path': 'M0 0L24 0L12 24Z', 'windingRule': 'NONZERO'}]}]}
    out = os.path.join(tmp, 'probe.pptx')
    report.clear()
    convert([tree], out, overrides={'Probe': [font]})
    prs = Presentation(out)
    xml = prs.slides[0].shapes._spTree.xml
    assert len(prs.slides) == 1 and report['groups'] == 1, report
    assert 'custGeom' in xml and 'roundRect' in xml and 'outerShdw' in xml and '>AA A<' in xml, 'missing shape kinds'
    assert prs.part._element.get('embedTrueTypeFonts') == '1', 'fonts not flagged as embedded'
    blob = next(p for p in prs.part.package.iter_parts() if str(p.partname).endswith('.fntdata')).blob
    assert struct.unpack('<L', blob[8:12])[0] == 0x00020001 and struct.unpack('<H', blob[34:36])[0] == 0x504C, 'bad EOT header'
    assert struct.unpack('<L', blob[:4])[0] == len(blob), 'EOT size field'
    print(f'self-test passed: {dict(report)}')


def main():
    for stream in (sys.stdout, sys.stderr):   # Windows pipes default to cp949/cp1252 and choke on frame names
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')
    ap = argparse.ArgumentParser(description='Figma frames -> fully editable PowerPoint (.pptx)')
    ap.add_argument('target', nargs='?', help='Figma URL (node-id of a page, section or frame) or a file key')
    ap.add_argument('-n', '--node', help='node id such as 280:2 (overrides the URL)')
    ap.add_argument('-o', '--out', help='output .pptx (default: the page or frame name)')
    ap.add_argument('--match', help='regex: only frames whose name matches')
    ap.add_argument('--sort', choices=('position', 'name'), default='position', help='slide order: canvas reading order or natural name order')
    ap.add_argument('--group-depth', type=int, default=3, help='Figma layer levels kept as PowerPoint groups (0 = none)')
    ap.add_argument('--font', action='append', default=[], metavar='FAMILY=FILE[,FILE]', help='font files for a family that is not on Google Fonts')
    ap.add_argument('--no-embed', action='store_true', help='reference fonts by name only (smaller file, looks right only where installed)')
    ap.add_argument('--full-fonts', action='store_true', help='embed whole fonts instead of the characters used (+ all Hangul)')
    ap.add_argument('--cache', help='folder to keep Figma node data between runs')
    ap.add_argument('--self-test', action='store_true', help='offline check of the pipeline')
    args = ap.parse_args()
    if args.self_test:
        return self_test()
    if not args.target:
        ap.error('a Figma URL or file key is required')
    run(args)


if __name__ == '__main__':
    main()
